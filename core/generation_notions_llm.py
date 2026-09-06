"""
Génération d'une structure de notions depuis un document (Partie 1,
06/09/2026 -- voir clovis-plan-travail-10-parties.md, Partie 1 : "la
fonctionnalité générer une structure depuis un document nécessite un
appel LLM dédié -- nouveau fichier séparé pour cette génération, ne pas
la mélanger avec le CRUD de base").

Ce fichier NE TOUCHE JAMAIS la base : il extrait le texte d'un document
et renvoie une structure proposée (arbre nom + enfants), à charge du
frontend (Partie 2) de l'afficher pour validation/édition avant que le
prof ne la crée réellement via les endpoints CRUD (api/programme_notions.py).

Extraction de texte dupliquée volontairement depuis api/uploads.py
(fonctions privées de ce fichier, pas importables proprement) et
core/bibliotheque_rag.py -- même principe déjà appliqué là
("dupliquée volontairement pour ne pas créer de dépendance croisée
entre les deux circuits").
"""

import io
import json
import logging

from constantes_agent import get_secret, GROQ_PRIMARY
from groq import Groq

logging.basicConfig(level=logging.INFO)

FORMATS_AUTORISES_GENERATION = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
}
TAILLE_MAX_DOCUMENT_GENERATION_OCTETS = 15 * 1024 * 1024  # 15 Mo, même limite que les documents de chat (api/uploads.py)
LONGUEUR_MAX_TEXTE_ENVOYE_AU_LLM = 30_000  # caractères, même garde-fou que api/uploads.py::LONGUEUR_MAX_TEXTE_EXTRAIT

PROFONDEUR_MAX_STRUCTURE = 4  # garde-fou : un document mal structuré ne doit pas produire un arbre infini


def _extraire_texte_pdf(contenu_bytes: bytes) -> str:
    import PyPDF2

    reader = PyPDF2.PdfReader(io.BytesIO(contenu_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extraire_texte_docx(contenu_bytes: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(contenu_bytes))
    morceaux = [p.text for p in document.paragraphs]
    for table in document.tables:
        for ligne in table.rows:
            morceaux.append("\t".join(cellule.text for cellule in ligne.cells))
    return "\n".join(morceaux)


def _extraire_texte_xlsx(contenu_bytes: bytes) -> str:
    import openpyxl

    classeur = openpyxl.load_workbook(io.BytesIO(contenu_bytes), data_only=True)
    morceaux = []
    for feuille in classeur.worksheets:
        morceaux.append(f"--- Feuille : {feuille.title} ---")
        for ligne in feuille.iter_rows(values_only=True):
            morceaux.append("\t".join("" if v is None else str(v) for v in ligne))
    return "\n".join(morceaux)


def _extraire_texte(contenu_bytes: bytes, content_type: str) -> str:
    extension = FORMATS_AUTORISES_GENERATION.get(content_type)
    if extension == "pdf":
        return _extraire_texte_pdf(contenu_bytes)
    if extension == "docx":
        return _extraire_texte_docx(contenu_bytes)
    if extension == "xlsx":
        return _extraire_texte_xlsx(contenu_bytes)
    raise ValueError(f"content_type non supporté : {content_type}")


_PROMPT_SYSTEME = (
    "Tu analyses un document de cours (programme, sommaire, table des "
    "matières, plan de chapitre...) et tu en extrais une structure de "
    "notions à enseigner, organisée en arbre (une notion peut avoir des "
    "sous-notions, jusqu'à {profondeur_max} niveaux de profondeur maximum). "
    "Réponds UNIQUEMENT en JSON, sous cette forme exacte :\n"
    '{{"notions": [{{"nom": "...", "enfants": [{{"nom": "...", "enfants": []}}]}}]}}\n'
    "Chaque \"nom\" est court et clair (le titre d'une notion ou sous-notion, "
    "pas une phrase). \"enfants\" est toujours présent, liste vide si la "
    "notion n'a pas de sous-notion. Ne réponds RIEN d'autre que ce JSON."
).format(profondeur_max=PROFONDEUR_MAX_STRUCTURE)


def _profondeur_reelle(noeuds: list[dict]) -> int:
    if not noeuds:
        return 0
    return 1 + max((_profondeur_reelle(n.get("enfants") or []) for n in noeuds), default=0)


def _nettoyer_arbre(noeuds, profondeur_restante: int) -> list[dict]:
    """Valide/nettoie la sortie du LLM : garde seulement nom (str non
    vide) + enfants (liste), tronque au-delà de PROFONDEUR_MAX_STRUCTURE
    plutôt que de faire échouer toute la génération pour un excès de
    profondeur isolé."""
    if not isinstance(noeuds, list):
        return []
    resultat = []
    for n in noeuds:
        if not isinstance(n, dict):
            continue
        nom = str(n.get("nom") or "").strip()
        if not nom:
            continue
        enfants = _nettoyer_arbre(n.get("enfants") or [], profondeur_restante - 1) if profondeur_restante > 1 else []
        resultat.append({"nom": nom, "enfants": enfants})
    return resultat


def generer_structure_notions_depuis_document(contenu_bytes: bytes, content_type: str) -> dict | None:
    """Renvoie {"notions": [...]} (arbre proposé) ou None en cas d'échec
    (extraction impossible, appel LLM en échec, ou réponse illisible) --
    l'appelant (api/programme_notions.py) traduit None en erreur API.
    Ne persiste RIEN en base : voir le docstring de ce fichier."""
    try:
        texte = _extraire_texte(contenu_bytes, content_type)
    except Exception as e:
        logging.error(f"ERREUR extraction texte (génération structure notions) : {e}")
        return None

    texte = (texte or "").strip()
    if not texte:
        return None
    texte = texte[:LONGUEUR_MAX_TEXTE_ENVOYE_AU_LLM]

    try:
        client = Groq(api_key=get_secret("GROQ_API_KEY"), max_retries=0)
        completion = client.chat.completions.create(
            model=GROQ_PRIMARY,
            messages=[
                {"role": "system", "content": _PROMPT_SYSTEME},
                {"role": "user", "content": texte},
            ],
            response_format={"type": "json_object"},
        )
        brut = json.loads(completion.choices[0].message.content or "{}")
    except Exception as e:
        logging.error(f"ERREUR Groq (génération structure notions) : {e}")
        return None

    notions = _nettoyer_arbre(brut.get("notions") or [], PROFONDEUR_MAX_STRUCTURE)
    if not notions:
        return None
    return {"notions": notions}
